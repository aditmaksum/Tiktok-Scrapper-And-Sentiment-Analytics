"""One-off: compare candidate sentiment models against the Tahap B labeled sample.

Not part of the runtime pipeline. Used once to decide which model (if any)
clears the G-08 accuracy gate before committing to it in model_classifier.py.
"""
from openpyxl import load_workbook
from transformers import pipeline

LABELED_XLSX = "runs/2026-08/labeling_sample_labeled.xlsx"

CANDIDATES = [
    {
        "name": "Aardiiiiy/indobertweet-base-Indonesian-sentiment-analysis",
        "label_map": {"negative": "negatif", "neutral": "netral", "positive": "positif"},
    },
    {
        "name": "mdhugol/indonesia-bert-sentiment-classification",
        "label_map": {"label_0": "positif", "label_1": "netral", "label_2": "negatif"},
    },
]


def load_labeled_sample():
    wb = load_workbook(LABELED_XLSX)
    ws = wb.active
    headers = [c.value for c in ws[1]]
    idx = {name: headers.index(name) for name in ("comment", "sentiment_label")}
    rows = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        label = row[idx["sentiment_label"]]
        if isinstance(label, str):
            label = label.strip().lower()
        rows.append({"comment": row[idx["comment"]] or "", "label": label})
    return rows


def main():
    rows = load_labeled_sample()
    print("loaded %d labeled comments\n" % len(rows))

    for candidate in CANDIDATES:
        print("=== %s ===" % candidate["name"])
        pipe = pipeline(
            "sentiment-analysis", model=candidate["name"], truncation=True, max_length=512
        )

        correct = 0
        for row in rows:
            text = row["comment"].strip()
            if not text:
                pred = "netral"
            else:
                raw_label = str(pipe(row["comment"])[0]["label"]).lower()
                pred = candidate["label_map"].get(raw_label, raw_label)
            if pred == row["label"]:
                correct += 1

        accuracy = correct / len(rows)
        print("accuracy: %.1f%% (%d/%d)\n" % (accuracy * 100, correct, len(rows)))


if __name__ == "__main__":
    main()
