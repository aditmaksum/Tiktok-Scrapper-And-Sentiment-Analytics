"""Tahap B calibration: measure the local model's accuracy against the
200-comment human-labeled sample, and pick ambiguous_confidence_threshold.

One-off dev tool (docs/plans/2026-08-30-model-cascade-implementation.md,
section 5). Not part of the runtime pipeline.

Usage:
    python scripts/calibrate_threshold.py

Reads:  local/labeling_sample_labeled.xlsx (comment_id, comment, sentiment_label, notes)
Writes: config/thresholds.yaml (ambiguous_confidence_threshold)
        config/calibration/tahap-b-labels.csv (comment_id+label+notes ONLY, no
        comment text/username - tracked answer key, see docs/calibration/codebook.md)
Prints: G-08 accuracy gate result + threshold sweep table
"""
import csv
import os

from openpyxl import load_workbook

from sosmed_sentiment.sentiment import model_classifier as mc

LABELED_XLSX = "local/labeling_sample_labeled.xlsx"
ACCURACY_TARGET = 0.80  # docs/plans/.../2026-08-30-model-cascade-implementation.md line 137
OUTPUT_CONFIG = "config/thresholds.yaml"
ANSWER_KEY_CSV = "config/calibration/tahap-b-labels.csv"

# Escalating everything below this ratio would blow the LLM budget; escalating
# nothing defeats the point of having an escalation layer at all. Sweep
# candidates and report cost/accuracy tradeoff at each rather than assume one.
THRESHOLD_CANDIDATES = [0.50, 0.60, 0.70, 0.80, 0.90, 0.95, 0.99]


def load_labeled_sample():
    wb = load_workbook(LABELED_XLSX)
    ws = wb.active
    headers = [c.value for c in ws[1]]
    idx = {name: headers.index(name) for name in ("comment_id", "comment", "sentiment_label", "notes")}

    rows = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        label = row[idx["sentiment_label"]]
        if isinstance(label, str):
            label = label.strip().lower()
        rows.append({
            "comment_id": row[idx["comment_id"]],
            "comment": row[idx["comment"]] or "",
            "label": label,
            "notes": row[idx["notes"]] or "",
        })
    return rows


def export_answer_key(rows):
    """Committed id+label+notes only - no comment text/username.

    This is the reproducible, git-trackable half of Tahap B labeling. The
    full file (local/labeling_sample_labeled.xlsx, with real customer
    comment text) stays gitignored - see .gitignore's local/ note.
    """
    os.makedirs(os.path.dirname(ANSWER_KEY_CSV), exist_ok=True)
    with open(ANSWER_KEY_CSV, "w", encoding="utf-8", newline="") as handle:
        writer = csv.writer(handle)
        writer.writerow(["comment_id", "sentiment_label", "notes"])
        for row in rows:
            writer.writerow([row["comment_id"], row["label"], row["notes"]])
    print("wrote %d rows to %s" % (len(rows), ANSWER_KEY_CSV))


def main():
    rows = load_labeled_sample()
    print("loaded %d labeled comments" % len(rows))

    export_answer_key(rows)

    mc.load_model()

    for row in rows:
        row["prediction"] = mc.classify(row["comment"])

    correct = sum(1 for r in rows if r["prediction"]["sentiment_label"] == r["label"])
    overall_accuracy = correct / len(rows)
    print("\n=== G-08 rejection gate ===")
    print("overall model accuracy (no escalation): %.1f%% (%d/%d), target >=%.0f%%" % (
        overall_accuracy * 100, correct, len(rows), ACCURACY_TARGET * 100
    ))
    verdict = "PASS" if overall_accuracy >= ACCURACY_TARGET else "FAIL"
    print("verdict: %s" % verdict)

    print("\n=== threshold sweep (escalate below threshold) ===")
    print("%-10s %-12s %-20s %-20s" % ("threshold", "escalated", "kept-accuracy", "escalated-accuracy"))
    best_threshold = None
    for threshold in THRESHOLD_CANDIDATES:
        kept = [r for r in rows if r["prediction"]["sentiment_confidence"] >= threshold]
        escalated = [r for r in rows if r["prediction"]["sentiment_confidence"] < threshold]

        kept_correct = sum(1 for r in kept if r["prediction"]["sentiment_label"] == r["label"])
        kept_acc = kept_correct / len(kept) if kept else float("nan")

        esc_correct = sum(1 for r in escalated if r["prediction"]["sentiment_label"] == r["label"])
        esc_acc = esc_correct / len(escalated) if escalated else float("nan")

        print("%-10.2f %-12s %-20s %-20s" % (
            threshold,
            "%d (%.0f%%)" % (len(escalated), 100 * len(escalated) / len(rows)),
            "%.1f%%" % (kept_acc * 100) if kept else "n/a",
            "%.1f%%" % (esc_acc * 100) if escalated else "n/a",
        ))

        # Pick the lowest threshold where the kept (non-escalated) slice already
        # clears the accuracy target - escalating past that spends LLM budget
        # without buying back correctness that isn't already there.
        if best_threshold is None and kept and kept_acc >= ACCURACY_TARGET:
            best_threshold = threshold

    if best_threshold is None:
        best_threshold = THRESHOLD_CANDIDATES[-1]
        print(
            "\nno swept threshold reaches a kept-accuracy >= target; "
            "falling back to the highest candidate (%.2f) - escalates the most" % best_threshold
        )

    print("\nselected ambiguous_confidence_threshold: %.2f" % best_threshold)

    with open(OUTPUT_CONFIG, "w", encoding="utf-8") as handle:
        handle.write(
            "# Generated by scripts/calibrate_threshold.py against the Tahap B\n"
            "# 200-comment labeled sample (local/labeling_sample_labeled.xlsx,\n"
            "# answer key mirrored to config/calibration/tahap-b-labels.csv).\n"
            "# Overall model accuracy measured: %.1f%% (target >=%.0f%%, G-08: %s).\n"
            "ambiguous_confidence_threshold: %.2f\n" % (
                overall_accuracy * 100, ACCURACY_TARGET * 100, verdict, best_threshold
            )
        )
    print("wrote %s" % OUTPUT_CONFIG)


if __name__ == "__main__":
    main()
